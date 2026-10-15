
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE = 10 * 1024 * 1024


class DocumentProcessingError(Exception):
    """Raised when a document cannot be processed."""


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentProcessingError(
            "Unsupported file type. Use PDF, DOCX, or TXT."
        )

    if not content:
        raise DocumentProcessingError("The uploaded file is empty.")

    if len(content) > MAX_FILE_SIZE:
        raise DocumentProcessingError(
            "The maximum allowed file size is 10 MB."
        )

    try:
        if extension == ".pdf":
            with fitz.open(stream=content, filetype="pdf") as pdf:
                text = "\n".join(
                    page.get_text() for page in pdf
                )

        elif extension == ".docx":
            document = Document(BytesIO(content))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            tables = [
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]

            text = "\n".join(paragraphs + tables)

        else:
            text = content.decode("utf-8-sig")

    except Exception as error:
        raise DocumentProcessingError(
            "Unable to extract text from this document."
        ) from error

    if not text.strip():
        raise DocumentProcessingError(
            "No readable text was found in the document."
        )

    return text.strip()
